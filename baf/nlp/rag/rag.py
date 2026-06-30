from __future__ import annotations

from baf.exceptions.logger import logger
import base64
import os
import tempfile
from typing import TYPE_CHECKING, Callable

from baf.core.message import Message, MessageType
from baf.nlp.llm.llm import LLM

if TYPE_CHECKING:
    from baf.core.agent import Agent
    from baf.core.file import File
    from baf.core.session import Session
    from baf.nlp.nlp_engine import NLPEngine

try:
    from langchain_community.document_loaders import PyPDFLoader
    from langchain_core.documents.base import Document
    from langchain_core.vectorstores.base import VectorStore, VectorStoreRetriever
    from langchain_text_splitters.base import TextSplitter
except ImportError:
    logger.warning("langchain dependencies in RAG could not be imported. You can install them from the "
                   "requirements/requirements-extras.txt file")


class RAGMessage:
    """The result of a RAG execution.

    Args:
        llm_name (str): the name of the LLM that generated the answer
        question (str): the original question
        answer (str): the generated answer
        docs (list[langchain_core.documents.base.Document]): the list of documents that were used as additional context
            for the answer generation

    Attributes:
        llm_name (str): the name of the LLM that generated the answer
        question (str): the original question
        answer (str): the generated answer
        docs (list[langchain_core.documents.base.Document]): the list of documents that were used as additional context
            for the answer generation
    """

    def __init__(self, llm_name: str, question: str, answer: str, docs: list[Document]):
        self.llm_name: str = llm_name
        self.question: str = question
        self.answer: str = answer
        self.docs: list[Document] = docs

    def to_dict(self):
        """Returns a dictionary containing the attributes of the RAGMessage object."""
        return {
            "llm_name": self.llm_name,
            "question": self.question,
            "answer": self.answer,
            "docs": [{
                'content': doc.page_content,
                'metadata': doc.metadata,
            } for doc in self.docs]
        }


class RAG:
    """A Retrieval Augmented Generation (RAG) implementation.

    A vector stores contains vectorized representations (i.e., embeddings) of chunks of data (text). For a given input
    query, a retriever gets the `k` most similar stored embeddings to the input embedding.

    This is usually used to, given a query or question, retrieve those chunks that could be helpful to give it an answer.
    Then, an LLM is in charge of generating that answer, given the original query and the retrieved data as context.

    Args:
        agent (Agent): the agent the RAG engine belongs to
        vector_store (langchain_core.vectorstores.base.VectorStore | Callable[[], VectorStore]): the vector store of the RAG engine, , or
            a zero-arg factory that produces one. A factory is recommended when ``session_scoped=True`` so that each
            session's RAG receives its own isolated store
        splitter (langchain_text_splitters.base.TextSplitter): the text splitter of the RAG engine
        llm_name (str): the name of the LLM of the RAG engine. It must have been previously created and assigned to the
            agent
        llm_prompt (str): the prompt containing the detailed instructions for the answer generation by the LLM. If none
            is provided, the :any:`default prompt <RAG.DEFAULT_LLM_PROMPT>` will be used
        k (int): number of chunks to retrieve from the vector store
        num_previous_messages (int): number of previous messages of the conversation to add to the LLM prompt context.
            Necessary a connection to :class:`~baf.db.monitoring_db.MonitoringDB`.
        session_scoped (bool): True if each session should have its own RAG engine and vector store.

    Attributes:
        _nlp_engine (NLPEngine): the NLPEngine that handles the NLP processes of the agent the RAG engine belongs to
        vector_store (langchain_core.vectorstores.base.VectorStore): the vector store of the RAG engine
        splitter (langchain_text_splitters.base.TextSplitter): the text splitter of the RAG engine
        llm_name (str): the name of the LLM of the RAG engine. It must have been previously created and assigned to the
            agent
        llm_prompt (str): the prompt containing the detailed instructions for the answer generation by the LLM. If none
            is provided, the :any:`default prompt <RAG.DEFAULT_LLM_PROMPT>` will be used
        k (int): number of chunks to retrieve from the vector store
        num_previous_messages (int): number of previous messages of the conversation to add to the LLM prompt context.
            Necessary a connection to :class:`~baf.db.monitoring_db.MonitoringDB`.
    """

    DEFAULT_LLM_PROMPT = "You are an assistant for question-answering tasks. Based on the previous messages in the conversation (if provided), and additional context retrieved from a database (if provided), answer the user question. If you don't know the answer, just say that you don't know. Note that if the question refers to a previous message, you may have to ignore the context since it is retrieved from the database based only on the question (the retrieval does not take into account the previous messages). Use three sentences maximum and keep the answer concise"
    SUPPORTED_FORMATS = ('pdf', 'docx', 'txt', 'md')

    def __init__(
            self,
            agent: 'Agent',
            vector_store: VectorStore | Callable[[], VectorStore],
            splitter: TextSplitter,
            llm_name: str,
            llm_prompt: str = None,
            k: int = 4,
            num_previous_messages: int = 0,
            session_scoped: bool = False
    ):
        self._nlp_engine: 'NLPEngine' = agent.nlp_engine
        if callable(vector_store) and not hasattr(vector_store, 'add_documents'):
            vector_store = vector_store()
        self.vector_store: VectorStore = vector_store
        self.splitter: TextSplitter = splitter
        self.llm_name: str = llm_name
        if not llm_prompt:
            llm_prompt = RAG.DEFAULT_LLM_PROMPT
        self.llm_prompt = llm_prompt
        self.k: int = k
        self.num_previous_messages: int = num_previous_messages
        if not session_scoped:
            self._nlp_engine._rag = self
    
    def load_documents_from_path(path: str, fmt: str) -> list['Document']:
        """Load raw (un-chunked) LangChain Documents from a file path.

        Dispatches to the appropriate loader based on ``fmt``. Used internally by
        :meth:`RAG.load_documents` and :meth:`RAG.add_file`.

        Args:
            path (str): path to the file on disk
            fmt (str): lowercase file format

        Returns:
            list[Document]: raw LangChain Documents (not yet chunked)

        Raises:
            ValueError: if ``fmt`` is not in `RAG.SUPPORTED_FORMATS`
            ImportError: if ``fmt`` is ``'docx'`` and python-docx is not installed
        """
        
        if fmt not in RAG.SUPPORTED_FORMATS:
            raise ValueError(f"Unsupported format '{fmt}'. Supported formats: {RAG.SUPPORTED_FORMATS}")
        if fmt == 'pdf':
            return PyPDFLoader(path).load()
        elif fmt == 'docx':
            try:
                from docx import Document as DocxDocument
            except ImportError:
                raise ImportError('python-docx is required for DOCX support. Install with: pip install python-docx')
            docx_doc = DocxDocument(path)
            text = '\n'.join(p.text for p in docx_doc.paragraphs if p.text.strip())
            return [Document(page_content=text, metadata={'source': os.path.basename(path)})]
        elif fmt in ('txt', 'md'):
            with open(path, 'r', encoding='utf-8') as f:
                text = f.read()
            return [Document(page_content=text, metadata={'source': os.path.basename(path)})]

    def load_pdfs(self, path: str) -> int:
        """Load PDF files from a given location into the RAG's vector store.

        Args:
            path (str): the path where the files are located

        Returns:
            int: the number of chunks added to the vector store
        """
        documents = []
        for file in os.listdir(path):
            if file.endswith('.pdf'):
                pdf_path = os.path.join(path, file)
                loader = PyPDFLoader(pdf_path)
                documents.extend(loader.load())
        chunked_documents = self.splitter.split_documents(documents)
        n_chunks = len(chunked_documents)
        self.vector_store.add_documents(chunked_documents)
        logger.info(f'[RAG] Added {n_chunks} chunks to RAG\'s vector store. Total: {len(self.vector_store.get()["documents"])}')
        return n_chunks
    
    def is_empty(self) -> bool:
        """Return ``True`` if the vector store contains no documents.

        Returns:
            bool: ``True`` if the corpus is empty
        """
        return len(self.vector_store.get()['ids']) == 0

    def clear(self) -> None:
        """Remove all documents from the vector store."""
        ids = self.vector_store.get()['ids']
        if ids:
            self.vector_store.delete(ids)
        logger.info('[RAG] Vector store cleared')

    def load_documents(self, path: str, formats: list[str] = None) -> int:
        """Load documents from a directory into the RAG's vector store.

        Supports multiple file formats. Use :meth:`load_pdfs` if only PDF loading is needed.

        Args:
            path (str): path to the directory containing the files
            formats (list[str]): restrict loading to a subset of formats, e.g. ``['pdf', 'docx']``.
                If ``None``, all `RAG.SUPPORTED_FORMATS` are loaded.

        Returns:
            int: the number of chunks added to the vector store
        """
        allowed = set(formats) if formats else set(RAG.SUPPORTED_FORMATS)
        documents = []
        for fname in os.listdir(path):
            fmt = os.path.splitext(fname)[1].lower().lstrip('.')
            if fmt in allowed:
                documents.extend(RAG.load_documents_from_path(os.path.join(path, fname), fmt))
        chunked_documents = self.splitter.split_documents(documents)
        n_chunks = len(chunked_documents)
        self.vector_store.add_documents(chunked_documents)
        logger.info(f'[RAG] Added {n_chunks} chunks to RAG\'s vector store. Total: {len(self.vector_store.get()["documents"])}')
        return n_chunks

    def add_file(self, file: 'File') -> int:
        """Extract text from a BAF :class:`~baf.core.file.File` object and add it to the RAG's vector store at runtime.

        Args:
            file (File): the BAF File object to load

        Returns:
            int: the number of chunks added to the vector store

        Raises:
            ValueError: if the file type is not in RAG.SUPPORTED_FORMATS`
        """
        fmt = (file.type or '').lower().lstrip('.')
        if not fmt or '/' in fmt:  # MIME type (e.g. "application/pdf") — derive from filename
            fmt = os.path.splitext(file.name or '')[1].lower().lstrip('.')
        file_bytes = base64.b64decode(file.base64)
        with tempfile.NamedTemporaryFile(suffix=f'.{fmt}', delete=False) as tmp:
            tmp.write(file_bytes)
            tmp_path = tmp.name
        try:
            documents = RAG.load_documents_from_path(tmp_path, fmt)
        finally:
            os.unlink(tmp_path)
        chunked_documents = self.splitter.split_documents(documents)
        n_chunks = len(chunked_documents)
        self.vector_store.add_documents(chunked_documents)
        logger.info(f'[RAG] Added {n_chunks} chunks from \'{file.name}\' to RAG\'s vector store.')
        return n_chunks

    def add_text(self, text: str, metadata: dict = None) -> int:
        """Chunk a raw text string and add it to the RAG's vector store at runtime.

        Args:
            text (str): the raw text to add
            metadata (dict): optional metadata to attach to each chunk

        Returns:
            int: the number of chunks added to the vector store
        """
        doc = Document(page_content=text, metadata=metadata or {})
        chunked_documents = self.splitter.split_documents([doc])
        n_chunks = len(chunked_documents)
        self.vector_store.add_documents(chunked_documents)
        logger.info(f'[RAG] Added {n_chunks} text chunks to RAG\'s vector store.')
        return n_chunks

    def run_retrieval(self, question: str, k: int = None) -> list[Document]:
        """Run retrieval. Given a query, return the `k` most relevant documents (i.e., chunks) from the RAG's vector store.

        Args:
            question (str): the input query
            k (int): the number of (top) documents to get. If none is provided, the RAG's default value will be used

        Returns:
            list[langchain_core.documents.base.Document]: the retrieved documents
        """
        if not k:
            k = self.k
        retriever: VectorStoreRetriever = self.vector_store.as_retriever(search_kwargs={"k": k}) # todo: set another attr for the retriever? (to be customizable)
        return retriever.invoke(question)

    def create_prompt(
            self,
            history: list[Message],
            docs: list[Document],
            question: str,
            llm_prompt: str = None,
    ) -> str:
        """
        Creates the prompt for the LLM answer generation.

        Args:
            history (list[Message]): the chat history
            docs (list[langchain_core.documents.base.Document]): the retrieved documents to use as context in the prompt
            question (str): the user question
            llm_prompt (str): the prompt containing the detailed instructions for the answer generation by the LLM. If
                none is provided, the RAG's default value will be used

        Returns:
            str: the LLM prompt
        """
        if not llm_prompt:
            llm_prompt = self.llm_prompt
        result = [llm_prompt]
        for message in history:
            if message.type == MessageType.RAG_ANSWER:
                result.append(f'Assistant: {message.content["answer"]}')
            elif message.type == MessageType.STR:
                result.append(f'{"User" if message.is_user else "Assistant"}: {message.content}')

        formatted_docs = "\n\n".join(doc.page_content for doc in docs)
        result.append(f'Context: {formatted_docs}')
        result.append(f'Question: {question}')
        result.append(f'Answer: ')
        return '\n\n'.join(result)

    def run(
            self,
            message: str,
            session: 'Session' = None,
            llm_prompt: str = None,
            llm_name: str = None,
            k: int = None,
            num_previous_messages: int = None
    ) -> RAGMessage:
        """Run the RAG engine.

        Args:
            session (Session): the session of the user that started this request. Must be provided if the chat history
                wants to be added as context to the LLM prompt.
            message (str): the message to be used as RAG query
            llm_prompt (str): the prompt containing the detailed instructions for the answer generation by the LLM. If
                none is provided, the RAG's default value will be used
            llm_name (str): the name of the LLM to use. If none is provided, the RAG's default value will be used
            k (int): the number of (top) documents to get. If none is provided, the RAG's default value will be used
            num_previous_messages (int): number of previous messages of the conversation to add to the LLM prompt
                context. If none is provided, the RAG's default value will be used. Necessary a connection to
                :class:`~baf.db.monitoring_db.MonitoringDB`.

        Returns:
            RAGMessage: the resulting RAG message
        """
        if not message and not session:
            raise ValueError('RAG Run: Must provide either a message or a session')
        if not llm_name:
            llm_name = self.llm_name
        if not num_previous_messages:
            num_previous_messages = self.num_previous_messages
        if session and num_previous_messages > 0:
            history = session.get_chat_history(n=num_previous_messages)
        else:
            history = []
        docs: list[Document] = self.run_retrieval(question=message, k=k)
        prompt = self.create_prompt(history=history, docs=docs, question=message, llm_prompt=llm_prompt)
        llm: LLM = self._nlp_engine._llms[llm_name]
        llm_response: str = llm.predict(prompt)
        return RAGMessage(llm_name=llm_name, question=message, answer=llm_response, docs=docs)
